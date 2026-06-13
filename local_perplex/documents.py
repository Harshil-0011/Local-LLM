import os
from pathlib import Path

class DocumentManager:
    def __init__(self, doc_dir: str = None):
        if doc_dir is None:
            doc_dir = os.getenv("LOCAL_PERPLEX_DOCUMENT_DIR")
            if doc_dir is None:
                project_root = Path(__file__).parent.parent
                doc_dir = project_root / "documents"
        else:
            doc_dir = Path(doc_dir)
        
        self.doc_dir = Path(doc_dir)
        self.doc_dir.mkdir(parents=True, exist_ok=True)

    def extract_document_text(self, file: str | Path) -> str:
        file = Path(file)
        ext = file.suffix.lower()
        if ext in [".txt", ".md"]:
            with open(file, "r", encoding="utf-8") as f:
                return f.read()
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(file)
            content = ""
            for i in range(min(5, len(reader.pages))):
                page_text = reader.pages[i].extract_text() or ""
                content += page_text + " "
            return content
        if ext == ".docx":
            import docx
            doc = docx.Document(file)
            return "\n".join([p.text for p in doc.paragraphs])
        if ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(file, data_only=True)
            ws = wb.active
            content = ""
            for row in ws.iter_rows(max_row=50, values_only=True):
                content += ",".join([str(c) for c in row if c is not None]) + "\n"
            return content
        if ext == ".csv":
            with open(file, "r", encoding="utf-8") as f:
                return f.read()
        return ""

    def get_local_context(self):
        context = []
        for file in self.doc_dir.glob("*"):
            ext = file.suffix.lower()
            try:
                content = self.extract_document_text(file)
            except Exception:
                continue
            if not content:
                continue
            if ext in [".txt", ".md"]:
                context.append(f"Local File: {file.name}\nContent: {content[:2500]}")
            elif ext == ".pdf":
                context.append(f"Local PDF: {file.name}\nContent: {content[:3000]}")
            elif ext == ".docx":
                context.append(f"Local Word Doc: {file.name}\nContent: {content[:3000]}")
            elif ext == ".xlsx":
                context.append(f"Local Excel (Data): {file.name}\nContent: {content[:4000]}")
            elif ext == ".csv":
                context.append(f"Local CSV (Data): {file.name}\nContent: {content[:4000]}")
        return "\n\n".join(context)

    def add_document(self, filename: str, content: str):
        safe_filename = os.path.basename(filename)
        with open(self.doc_dir / safe_filename, "w", encoding="utf-8") as f:
            f.write(content)
